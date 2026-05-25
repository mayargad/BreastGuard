"""
Lab Report File Parser
Extracts blood marker values from uploaded PDF or Word (.docx) files
using regex pattern matching on the extracted text.
"""
import re
import io
from flask import Blueprint, request, jsonify

parse_bp = Blueprint('parse', __name__)

# ── Marker extraction patterns ─────────────────────────────────
# Supports formats: "CA 15-3: 25.4", "CA 15-3  25.4", "CA15-3=25.4"
# Also handles comma as decimal separator (e.g. 25,4)
MARKER_PATTERNS = {
    'ca153': [
        r'CA\s*15[.\s\-\u2013\u2014]*3\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'Cancer\s*Antigen\s*15[.\s\-\u2013\u2014]*3\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
    ],
    'cea': [
        r'CEA\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'C\.E\.A\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'Carcinoembryonic\s*Antigen\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
    ],
    'wbc': [
        r'WBC\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'White\s*Blood\s*Cells?\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'Leukocytes?\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'TLC\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
    ],
    'rbc': [
        r'RBC\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'Red\s*Blood\s*Cells?\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'Erythrocytes?\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
    ],
    'hgb': [
        r'\bHGB\b\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'\bHgb\b\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'\bHb\b\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'\bHB\b\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'H(?:a?e?mo|emo)globin\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
    ],
    'plt': [
        r'\bPLT\b\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'Platelets?\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'Thrombocytes?\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
        r'Plt\s*Count\s*[:\=\s|]+([0-9]+[.,][0-9]+|[0-9]+)',
    ],
}

# Sanity ranges — ignore wildly out-of-range extractions (OCR errors etc.)
SANITY_RANGES = {
    'ca153': (0, 5000),
    'cea':   (0, 1000),
    'wbc':   (0.1, 100),
    'rbc':   (0.5, 10),
    'hgb':   (1, 25),
    'plt':   (1, 2000),
}


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract all text from a PDF file using pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text_parts.append(t)
        return '\n'.join(text_parts)
    except Exception as e:
        raise ValueError(f"Could not read PDF: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract all text from a Word (.docx) file, preserving table row structure."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        lines = []

        # Paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text.strip())

        # Tables — join each row's cells with ' | ' to keep name/value together
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    lines.append(' | '.join(cells))

        return '\n'.join(lines)
    except Exception as e:
        raise ValueError(f"Could not read Word file: {str(e)}")


def normalize_value(raw: str) -> float:
    """Convert extracted string to float, handling comma decimals (e.g. '25,4' → 25.4)."""
    return float(raw.replace(',', '.'))


def parse_markers_from_text(text: str) -> dict:
    """
    Run all regex patterns over the extracted text and return
    a dict of {marker_id: value_string} for found markers.
    """
    found = {}

    # Normalize: collapse multiple spaces/tabs, keep newlines for context
    text = re.sub(r'[ \t]+', ' ', text)

    for marker_id, patterns in MARKER_PATTERNS.items():
        lo, hi = SANITY_RANGES[marker_id]
        for pattern in patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            for m in matches:
                try:
                    val = normalize_value(m)
                    if lo <= val <= hi:
                        found[marker_id] = str(round(val, 2))
                        break
                except ValueError:
                    continue
            if marker_id in found:
                break

    return found


@parse_bp.route('/extract', methods=['POST'])
def extract_from_file():
    """
    POST endpoint — accepts multipart/form-data with a 'file' field.
    Returns extracted marker values as JSON.
    """
    if 'file' not in request.files:
        return jsonify({"error": "No file uploaded"}), 400

    f = request.files['file']
    filename = f.filename.lower()

    if not filename:
        return jsonify({"error": "Empty filename"}), 400

    file_bytes = f.read()
    if len(file_bytes) == 0:
        return jsonify({"error": "Uploaded file is empty"}), 400

    # ── Extract text ───────────────────────────────────────────
    try:
        if filename.endswith('.pdf'):
            text = extract_text_from_pdf(file_bytes)
        elif filename.endswith(('.docx', '.doc')):
            text = extract_text_from_docx(file_bytes)
        else:
            return jsonify({"error": "Unsupported file type. Please upload PDF or DOCX."}), 400
    except ValueError as e:
        return jsonify({"error": str(e)}), 422

    if not text.strip():
        return jsonify({
            "markers": {},
            "found_count": 0,
            "warning_en": "Could not extract readable text from this file. It may be a scanned image — please enter values manually.",
            "warning_ar": "تعذّر استخراج نص من هذا الملف. قد يكون صورة ممسوحة ضوئياً — يرجى إدخال القيم يدوياً."
        }), 200

    # ── Parse markers ──────────────────────────────────────────
    markers = parse_markers_from_text(text)

    return jsonify({
        "markers": markers,
        "found_count": len(markers),
        "extracted_text_preview": text[:300] + ('...' if len(text) > 300 else ''),
        "success_en": f"Successfully extracted {len(markers)} marker(s) from your file." if markers else "No recognized markers found. You can enter values manually.",
        "success_ar": f"تم استخراج {len(markers)} مؤشر(ات) من ملفك بنجاح." if markers else "لم يتم العثور على مؤشرات معروفة. يمكنك إدخال القيم يدوياً."
    }), 200